"""配布 Web UI（src/hoiku_agent/web）の決定論スモーク（LLM 非依存・creds 不要）。

生成そのもの（/run_sse 経由の日誌/月案）は LLM が要るのでここでは検証しない（層B eval / 実機スモーク）。
ここで担保するのは「配線が崩れていないこと」＝静的配信・自前 API の形・利用枠の応答・/ の着地。
"""

from __future__ import annotations

from datetime import datetime
import re

import pytest
import server
from fastapi.testclient import TestClient
from hoiku_agent.config import settings
from hoiku_agent.harness import demo_seed, record_store
from hoiku_agent.harness.child_record_period import parse_child_record_period

# autouse スタブから実物へ戻すための参照（初回ログイン auto-seed の実挙動テスト用）。
_REAL_SEED_WORKSPACE = demo_seed.seed_workspace


def _client() -> TestClient:
    # follow_redirects=False で / → /app/ のリダイレクトを検証できるようにする。
    return TestClient(server.app, follow_redirects=False)


def _sign_in_with_google(c: TestClient, monkeypatch, *, email: str = "sensei@example.com") -> None:
    """公式ボタンの POST を模して、検証済み Google session を作る（外部通信なし）。"""
    from hoiku_agent.web import auth

    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    monkeypatch.setattr(
        auth,
        "validate_google_credential",
        lambda credential: auth.GoogleUser(subject="google-subject-123", email=email),
    )
    welcome = c.get("/?next=/app/")  # popup callback 用の戻り先と session CSRF token を作る。
    csrf = re.search(r'"X-Login-CSRF": "([^"]+)"', welcome.text).group(1)
    r = c.post(
        "/auth/google",
        json={"credential": "signed-token"},
        headers={"Origin": "http://testserver", "X-Login-CSRF": csrf},
    )
    assert r.status_code == 200 and r.json()["redirect"] == "/app/"


@pytest.fixture(autouse=True)
def _stub_auto_seed(monkeypatch):
    """初回ログインのデフォルト seed を既定で止める（呼び出しは記録）。

    サインイン＋DB を併用する既存テストの件数前提（児童・書類が空から始まる）を守る。
    auto-seed / 初期化の実挙動を検証するテストだけ `real_seed` fixture で実物へ戻す。
    """
    calls: list[tuple] = []

    def _stub(workspace_id, **kwargs):
        calls.append((workspace_id, kwargs))
        return {"status": "skipped", "reason": "test stub"}

    monkeypatch.setattr(demo_seed, "seed_workspace", _stub)
    yield calls


@pytest.fixture()
def real_seed(monkeypatch):
    """auto-seed スタブを実物へ戻す（sqlite へ実際に seed する検証用）。"""
    monkeypatch.setattr(demo_seed, "seed_workspace", _REAL_SEED_WORKSPACE)


def test_config_shape() -> None:
    r = _client().get("/api/config")
    assert r.status_code == 200
    body = r.json()
    assert body["app_name"] == "hoiku_agent"
    assert body["default_user_id"] == "caregiver"
    for key in ("memory_connected", "rag_connected", "llm_budget", "model"):
        assert key in body
    assert parse_child_record_period(body["current_child_record_period"]) is not None
    assert body["current_child_record_period"] in {
        option["value"] for option in body["child_record_periods"]
    }
    assert all(
        {"value", "label", "fiscal_year", "quarter"} <= set(option)
        for option in body["child_record_periods"]
    )
    assert {"available", "limit_yen", "used_yen", "remaining_yen"} <= body["llm_budget"].keys()
    # 園の実 Word 様式に対応済みの kind を UI の出し分け用に返す（保育経過記録・クラス月案は配線済み）。
    assert "child_record" in body["docx_kinds"]
    assert "class_monthly" in body["docx_kinds"]
    # レビュー巡回の上限＝harness の SSOT を露出（UI が差し戻し時「N巡目/最大M」の M に使う）。
    from hoiku_agent.harness.pipeline import MAX_REVIEW_ITERATIONS

    assert body["max_review_iterations"] == MAX_REVIEW_ITERATIONS


def test_doc_template_shape() -> None:
    """/api/doc-template＝編集フォームが本文の順序/ラベルに使う（4種別・key/label/kind）。"""
    body = _client().get("/api/doc-template").json()
    assert set(body["templates"]) == {"diary", "monthly", "child_record", "nursery_record"}
    diary = body["templates"]["diary"]
    assert diary and diary[0]["key"] == "daily_aim"
    for sec in diary:
        assert {"key", "label", "kind"} <= set(sec)


def test_static_ui_served() -> None:
    c = _client()
    # SPA 本体と各 ES モジュールが配信される。
    page = c.get("/app/")
    assert page.status_code == 200
    assert "データを初期化して始める" in page.text
    assert 'id="record-period"' in page.text
    assert 'id="record-start"' not in page.text
    assert 'id="record-end"' not in page.text
    for asset in (
        "app.js",
        "adk.js",
        "docflow.js",
        "docedit.js",
        "policy.js",
        "notation.js",
        "records.js",
        "classes.js",
        "diaryform.js",
        "ui.js",
        "styles.css",
    ):
        assert c.get(f"/app/{asset}").status_code == 200, asset

    # 初期化成功後は workspace 単位の完了印を保存し、再読込後も一回限りのボタンを出さない。
    app_js = c.get("/app/app.js").text
    assert 'const DATA_INITIALIZED_KEY_PREFIX = "hoiku_data_initialized:"' in app_js
    assert "rememberDataInitialized(cfg);" in app_js
    assert 'resetData.classList.add("hidden");' in app_js


def test_edit_textareas_grow_with_content_without_inner_scroll() -> None:
    """編集欄は初期表示・入力の両方で内容高に追従し、欄内スクロールを作らない。"""
    c = _client()
    script = c.get("/app/docedit.js").text
    styles = c.get("/app/styles.css").text

    assert 't.style.height = "auto"' in script
    assert "t.scrollHeight + borderHeight" in script
    assert 't.addEventListener("input"' in script
    assert "requestAnimationFrame" in script
    assert "new ResizeObserver" in script
    assert "textarea.de-input{overflow-y:hidden;resize:none}" in styles


def test_root_shows_welcome() -> None:
    # 配布リンクの素の URLは、強制遷移せず案内画面を表示する。
    r = _client().get("/")
    assert r.status_code == 200
    assert "保育の記録を" in r.text
    assert "/public/welcome.css" in r.text


def test_policy_route_returns_cards_and_history() -> None:
    """/api/policy は育つ指針＝構造化カード＋変更履歴＋store を返す（閲覧・素通し）。"""
    c = _client()
    r = c.get("/api/policy")
    assert r.status_code == 200
    body = r.json()
    assert isinstance(body.get("cards"), list)
    assert isinstance(body.get("history"), list)
    assert body.get("store") in ("persistent", "ephemeral", "unavailable")
    # seed（共通/月案）が読める環境ではカードが入る。カード形（doc_type/body）も確認。
    # 参照方針も他のカードと同じ自然文の1枚として入る（専用の kind/references は持たない・2026-07-12簡素化）。
    if body["cards"]:
        card = body["cards"][0]
        assert {"id", "body", "doc_type", "doc_label"} <= card.keys()
        assert "kind" not in card and "references" not in card


def test_list_apps_has_root_agent() -> None:
    assert "hoiku_agent" in _client().get("/list-apps").json()


def test_llm_budget_requires_google_login_when_enabled(monkeypatch) -> None:
    """本番相当では作成・改善・取込・校正のいずれも Google ログインなしに開かない。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    assert c.post("/run_sse", json={}).status_code == 401
    assert c.post("/api/improve", json={"diff": "x"}).status_code == 401
    assert c.get("/api/config").status_code == 401


def test_improver_resume_rejects_different_workspace(monkeypatch) -> None:
    """保留中の改善セッションは開始時と異なる workspace から再開できない。"""
    from hoiku_agent.web import improver_stream

    sid = "workspace-bound-session"
    improver_stream._SESSIONS[sid] = (object(), "adk-session", "owner-workspace")
    monkeypatch.setattr(improver_stream, "resolve_workspace_id", lambda request, now: None)
    try:
        response = _client().post(
            "/api/improve/resume",
            json={"session_id": sid, "function_call_id": "ask-1", "answer": "反映する"},
        )
    finally:
        improver_stream._SESSIONS.pop(sid, None)

    assert response.status_code == 200
    assert "この改善セッションを再開する権限がありません" in response.text


def test_llm_budget_limit_returns_clear_message(monkeypatch) -> None:
    from hoiku_agent.harness import llm_budget
    from hoiku_agent.harness.llm_budget import BudgetDecision

    c = _client()
    _sign_in_with_google(c, monkeypatch)
    monkeypatch.setattr(
        llm_budget,
        "reserve",
        lambda subject, path: BudgetDecision(False, "user_hourly_limit", 35),
    )
    r = c.post("/api/proofread", json={"kind": "diary", "entry": {}})
    assert r.status_code == 429
    assert r.json()["code"] == "user_hourly_limit"
    assert "AI利用枠に達しました" in r.json()["error"]


# ──────────────────── 編集UI backend（form-meta / finalize-edit・harness 中継） ────────────────────


def _edit_diary_entry() -> dict:
    """編集フォームが送る DiaryEntry 相当の dict（型を通す good 例）。"""
    return {
        "date": "2026-06-25",
        "age_band": "0-2",
        "weather": "晴れ",
        "daily_aim": "安心して好きな遊びに関わる",
        "attendance": [{"child_id": "架空児A", "present": True, "reason": None}],
        "health_notes": None,
        "practice_record": "園庭で砂遊びを行った。",
        "individual_notes": [
            {
                "child_id": "架空児A",
                "observed_state": "砂の感触を確かめた",
                "tags": ["身近なものと関わり感性が育つ"],
                "life_record": {"meal": "完食", "sleep": "午睡2時間"},
            }
        ],
        "evaluation": {"child_focus": "感触に集中", "self_review": "道具が適切"},
        "parent_contact": None,
    }


def test_form_meta_exposes_enum_vocab() -> None:
    """タグ選択肢は schemas の Enum が SSOT（告示準拠の正しい文言を含む）。"""
    body = _client().get("/api/form-meta").json()
    assert "健やかに伸び伸びと育つ" in body["three_viewpoint"]  # 告示準拠（「と」入り）
    assert body["five_domains"] == ["健康", "人間関係", "環境", "言葉", "表現"]
    assert "数量や図形、標識や文字などへの関心・感覚" in body["ten_no_sugata"]


def test_finalize_edit_diary_revalidates_and_formats() -> None:
    """編集後の dict を harness で再検査・再整形して返す（標準様式の見出し・型成立）。"""
    r = _client().post("/api/finalize-edit", json={"kind": "diary", "entry": _edit_diary_entry()})
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert body["problems"] == []
    assert "主な活動" in body["formatted"] and "個別の記録" in body["formatted"]


def test_finalize_edit_surfaces_validation_after_edit() -> None:
    """編集で生活記録を空にしたら不足を返す（編集後も型成立ゲートが効く）。"""
    entry = _edit_diary_entry()
    entry["individual_notes"][0]["life_record"] = {
        "meal": "",
        "sleep": "",
        "toilet": "",
        "mood_health": "",
    }
    body = _client().post("/api/finalize-edit", json={"kind": "diary", "entry": entry}).json()
    assert body["ok"] is False
    assert any("生活記録" in p for p in body["problems"])


def test_finalize_edit_class_monthly_revalidates_and_formats() -> None:
    """クラス月案の編集後 dict も harness で再検査・再整形できる（kind=class_monthly・§18）。"""
    r = _client().post(
        "/api/finalize-edit",
        json={"kind": "class_monthly", "entry": _edit_class_monthly_entry()},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert "月間指導計画" in body["formatted"] and "指導計画（区分×領域）" in body["formatted"]


def test_finalize_edit_class_monthly_surfaces_validation() -> None:
    """0–2 で個人目標を空にしたら不足を返す（編集後も型成立ゲートが効く）。"""
    entry = _edit_class_monthly_entry()
    entry["individual_goals"] = []
    body = (
        _client().post("/api/finalize-edit", json={"kind": "class_monthly", "entry": entry}).json()
    )
    assert body["ok"] is False
    assert any("個人目標" in p for p in body["problems"])


def _edit_child_record_entry() -> dict:
    """保育経過記録の編集フォーム相当 dict（型を通す good 例・§19）。"""
    return {
        "period": "2026-04〜2026-06",
        "age_band": "0-2",
        "child_id": "架空児A",
        "age_months": "1歳3か月",
        "development_notes": [
            {
                "description": "伝い歩きから一人歩きへ移行し、探索範囲が広がった",
                "tags": ["健やかに伸び伸びと育つ"],
            }
        ],
        "care_notes": "",
        "family_liaison": "連絡帳で歩行の様子を共有した",
        "overall_note": "安心できる関係を土台に自分から環境へ関わる姿が増えた",
        "next_aims": "",
    }


def test_finalize_edit_child_record_revalidates_and_formats() -> None:
    """保育経過記録の編集後 dict も harness で再検査・再整形できる（kind=child_record）。"""
    r = _client().post(
        "/api/finalize-edit", json={"kind": "child_record", "entry": _edit_child_record_entry()}
    )
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert "保育経過記録" in body["formatted"] and "総合所見" in body["formatted"]


def test_finalize_edit_child_record_surfaces_validation() -> None:
    """総合所見を空にしたら不足を返す（保育経過記録でも編集後の型成立ゲートが効く）。"""
    entry = _edit_child_record_entry()
    entry["overall_note"] = ""
    body = (
        _client().post("/api/finalize-edit", json={"kind": "child_record", "entry": entry}).json()
    )
    assert body["ok"] is False
    assert any("総合所見" in p for p in body["problems"])


def _edit_nursery_record_entry() -> dict:
    """保育要録の編集フォーム相当 dict（型を通す good 例・§19・L4）。"""
    return {
        "fiscal_year": "2026",
        "age_band": "3-5",
        "child_id": "架空児A",
        "age_months": "6歳0か月",
        "final_year_focus": "共通の目的に向かって思いや考えを出し合いながら活動を楽しむ",
        "individual_focus": "自分を発揮しながらさまざまな活動を楽しむ",
        "development_notes": [
            {"description": "友だちと考えを出し合い協力する姿が増えた", "tags": ["人間関係"]}
        ],
        "special_notes": "",
        "growth_until_final": "入園当初は不安が大きかったが、生き生きと表現を楽しむ姿へ育った",
    }


def test_finalize_edit_nursery_record_revalidates_and_formats() -> None:
    """保育要録の編集後 dict も harness で再検査・再整形できる（kind=nursery_record）。"""
    r = _client().post(
        "/api/finalize-edit",
        json={"kind": "nursery_record", "entry": _edit_nursery_record_entry()},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    assert "保育所児童保育要録" in body["formatted"]
    assert "最終年度に至るまでの育ち" in body["formatted"]


def test_finalize_edit_nursery_record_surfaces_validation() -> None:
    """最終年度に至るまでの育ちを空にしたら不足を返す（要録でも編集後の型成立ゲートが効く）。"""
    entry = _edit_nursery_record_entry()
    entry["growth_until_final"] = ""
    body = (
        _client().post("/api/finalize-edit", json={"kind": "nursery_record", "entry": entry}).json()
    )
    assert body["ok"] is False
    assert any("最終年度に至るまでの育ち" in p for p in body["problems"])


def test_finalize_edit_does_not_consume_llm_budget() -> None:
    """編集の再確定は LLM 非課金なので、利用枠の対象にしない。"""
    c = _client()
    assert c.get("/api/form-meta").status_code == 200
    r = c.post("/api/finalize-edit", json={"kind": "diary", "entry": _edit_diary_entry()})
    assert r.status_code == 200 and r.json()["ok"] is True


def test_finalize_edit_unknown_kind_is_400() -> None:
    """未知 kind は黙って diary として解釈せず 400（他の web 口と同じ kind 検証・黙って誤解釈しない）。"""
    r = _client().post(
        "/api/finalize-edit", json={"kind": "child-record", "entry": _edit_diary_entry()}
    )
    assert r.status_code == 400 and r.json()["code"] == "invalid_request"


def test_run_live_websocket_route_is_removed() -> None:
    """WebSocket /run_live は撤去済み（HTTP ミドルウェアのゲートを素通りする課金口を塞ぐ）。"""
    ws = [
        getattr(r, "path", None)
        for r in server.app.router.routes
        if type(r).__name__ in ("WebSocketRoute", "APIWebSocketRoute")
    ]
    assert "/run_live" not in ws


def test_dev_builder_memory_writes_are_forbidden_reads_open() -> None:
    """web=True が露出する ADK の dev/builder/memory 面：書込/実行（非 GET）は禁止、読取 GET は素通し。

    公開デモで未認証の LLM 課金迂回（/dev の eval 実行）・Memory Bank 直接書込（承認ゲート迂回）・
    エージェント定義タンパリング（/builder 書込）を塞ぐ。承認 PATCH・セッション作成は従来どおり開放。
    """
    from hoiku_agent.web.routes import _is_forbidden_dev_write

    # 書込/実行（非 GET）＝配布 UI では禁止
    assert _is_forbidden_dev_write("/builder/save", "POST") is True
    assert _is_forbidden_dev_write("/dev/apps/x/builder/save", "POST") is True
    assert _is_forbidden_dev_write("/dev/apps/x/eval_sets/y/run_eval", "POST") is True
    assert _is_forbidden_dev_write("/dev/apps/x/tests/run", "POST") is True
    assert _is_forbidden_dev_write("/apps/x/users/u/memory", "PATCH") is True
    # 読取・承認・セッション作成＝素通し
    assert _is_forbidden_dev_write("/builder/app/x", "GET") is False
    assert _is_forbidden_dev_write("/dev/apps/x/eval-results", "GET") is False
    assert _is_forbidden_dev_write("/apps/x/users/u/sessions/s", "PATCH") is False
    assert _is_forbidden_dev_write("/apps/x/users/u/sessions", "POST") is False


# ──────────────────── 帳票PDF 出力（/api/export-pdf・現場でそのまま綴じる最終形） ────────────────────


def _edit_monthly_entry() -> dict:
    """月案の編集フォーム相当 dict（帳票PDF 描画の入力）。"""
    return {
        "month": "2026-07",
        "age_band": "0-2",
        "child_id": "架空児A",
        "age_months": "1歳3か月",
        "prev_child_state": "前月は感触遊びに集中していた",
        "nurturing_life": "安全と生理的欲求の充足",
        "nurturing_emotion": "応答的関わりで安心を支える",
        "education": [{"aim": "素材に親しむ", "tags": ["身近なものと関わり感性が育つ"]}],
        "monthly_goals": "夏の遊びを楽しむ",
        "environment_support": "水遊びの環境を整える",
        "events_family_food": None,
        "evaluation_reflection": "翌月へつなげる",
    }


def _edit_class_monthly_entry() -> dict:
    """クラス月案の編集フォーム相当 dict（園の実様式＝区分×領域グリッド＋0–2 の個人目標）。"""
    grid = [
        {
            "category": cat,
            "domain": dom,
            "aim": f"{dom}のねらい",
            "environment": "環境",
            "child_state": "姿",
            "support": "配慮",
        }
        for cat, dom in (
            ("養護", "生命の保持"),
            ("養護", "情緒の安定"),
            ("教育", "健康"),
            ("教育", "人間関係"),
            ("教育", "環境"),
            ("教育", "言葉"),
            ("教育", "表現"),
        )
    ]
    return {
        "month": "2026-07",
        "age_band": "0-2",
        "class_name": "ひよこ組",
        "monthly_goal": "梅雨期も健康に過ごす",
        "prev_month_state": "前月は感触遊びに集中していた",
        "events": "七夕",
        "parent_support": "連絡帳で連携",
        "grid": grid,
        "syokuiku": "手づかみ食べ",
        "health_safety": "ブレスチェック",
        "family_liaison": "連絡帳",
        "staff_liaison": "申し送り",
        "individual_goals": [
            {
                "child_id": "架空児A",
                "age_months": "1歳3か月",
                "child_state": "歩行安定",
                "aim_support": "探索保障",
            }
        ],
    }


def _assert_is_pdf(r) -> None:
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"
    # 日本語ファイル名を RFC5987 で載せる（ダウンロード名が化けない）。
    assert "filename*=UTF-8''" in r.headers.get("content-disposition", "")


def test_export_pdf_diary_returns_pdf() -> None:
    r = _client().post("/api/export-pdf", json={"kind": "diary", "entry": _edit_diary_entry()})
    _assert_is_pdf(r)


def test_export_pdf_monthly_returns_pdf() -> None:
    r = _client().post("/api/export-pdf", json={"kind": "monthly", "entry": _edit_monthly_entry()})
    _assert_is_pdf(r)


def test_export_pdf_class_monthly_returns_pdf() -> None:
    r = _client().post(
        "/api/export-pdf", json={"kind": "class_monthly", "entry": _edit_class_monthly_entry()}
    )
    _assert_is_pdf(r)


def test_export_pdf_child_record_returns_pdf() -> None:
    r = _client().post(
        "/api/export-pdf", json={"kind": "child_record", "entry": _edit_child_record_entry()}
    )
    _assert_is_pdf(r)


def test_export_pdf_nursery_record_returns_pdf() -> None:
    r = _client().post(
        "/api/export-pdf", json={"kind": "nursery_record", "entry": _edit_nursery_record_entry()}
    )
    _assert_is_pdf(r)


def test_export_pdf_sparse_entry_still_renders() -> None:
    """空欄多めの entry でも帳票は出る（型検査はしない＝描画のみ・型の保証は harness）。"""
    r = _client().post(
        "/api/export-pdf",
        json={"kind": "diary", "entry": {"age_band": "0-2", "individual_notes": [{}]}},
    )
    _assert_is_pdf(r)


def test_export_pdf_invalid_kind_400() -> None:
    r = _client().post("/api/export-pdf", json={"kind": "weekly", "entry": _edit_diary_entry()})
    assert r.status_code == 400
    assert r.json()["code"] == "invalid_request"


def test_export_pdf_malformed_entry_is_400_not_500() -> None:
    """entry の内部が不正（リスト要素が dict でない）でも 500 でなく 400（公開口の契約「entry 不正は 400」）。"""
    r = _client().post(
        "/api/export-pdf",
        json={"kind": "diary", "entry": {"age_band": "0-2", "individual_notes": ["x"]}},
    )
    assert r.status_code == 400 and r.json()["code"] == "invalid_request"


def test_export_pdf_child_record_year_0000_period_not_500() -> None:
    """period に年 "0000" が入っても月齢自動充填で 500 にせず帳票は出る（_period_end_date が None 降格）。"""
    r = _client().post(
        "/api/export-pdf",
        json={
            "kind": "child_record",
            "entry": {
                "child_id": "架空児A",
                "period": "0000-04〜0000-06",
                "age_band": "0-2",
                "development_notes": [{"description": "a", "tags": ["健康な心と体"]}],
                "overall_note": "b",
            },
        },
    )
    _assert_is_pdf(r)


def test_export_pdf_very_long_body_still_renders() -> None:
    """1欄の本文が1ページ高を超えても splitInRow で次ページへ流し 500 にしない（LayoutError 回避）。"""
    long = "長い叙述である。" * 1500
    r = _client().post(
        "/api/export-pdf",
        json={
            "kind": "child_record",
            "entry": {
                "child_id": "架空児A",
                "period": "2026-04〜2026-06",
                "age_band": "0-2",
                "development_notes": [{"description": long, "tags": ["健康な心と体"]}],
                "overall_note": long,
            },
        },
    )
    _assert_is_pdf(r)


# ──────────── 園の実 Word 様式（/api/export-docx・Word 編集用の最終形・python-docx） ────────────


def _assert_is_docx(r) -> None:
    assert r.status_code == 200
    assert "wordprocessingml.document" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # docx = zip
    assert "filename*=UTF-8''" in r.headers.get("content-disposition", "")


def test_export_docx_child_record_returns_docx() -> None:
    r = _client().post(
        "/api/export-docx", json={"kind": "child_record", "entry": _edit_child_record_entry()}
    )
    _assert_is_docx(r)


def test_export_docx_monthly_returns_docx() -> None:
    r = _client().post("/api/export-docx", json={"kind": "monthly", "entry": _edit_monthly_entry()})
    _assert_is_docx(r)


def test_export_docx_class_monthly_returns_docx() -> None:
    r = _client().post(
        "/api/export-docx", json={"kind": "class_monthly", "entry": _edit_class_monthly_entry()}
    )
    _assert_is_docx(r)


def test_export_docx_nursery_record_returns_docx() -> None:
    r = _client().post(
        "/api/export-docx", json={"kind": "nursery_record", "entry": _edit_nursery_record_entry()}
    )
    _assert_is_docx(r)


def test_export_docx_unsupported_kind_400() -> None:
    """docx 未対応の kind（現状 diary 等）は 400 で正直に返す（握りつぶさない）。"""
    r = _client().post("/api/export-docx", json={"kind": "diary", "entry": _edit_diary_entry()})
    assert r.status_code == 400
    assert r.json()["code"] == "invalid_request"


def test_export_pdf_is_not_llm_budgeted() -> None:
    """帳票PDF 出力は LLM 非課金なので利用枠を消費しない。"""
    r = _client().post("/api/export-pdf", json={"kind": "diary", "entry": _edit_diary_entry()})
    _assert_is_pdf(r)


# ──────────────────── 書類アーカイブ（/api/records・/api/children＝record_store 中継・Phase 1） ────────────────────


@pytest.fixture()
def records_db(tmp_path, monkeypatch):
    """sqlite の一時 DB に向けてアーカイブのスキーマを作る（web 経由の決定論検証）。"""
    monkeypatch.setattr(settings, "database_url", f"sqlite:///{tmp_path}/archive.db")
    record_store.reset_engine_cache()
    record_store.Base.metadata.create_all(record_store._engine())
    yield
    record_store.reset_engine_cache()


def test_archive_write_maps_missing_schema_to_friendly_error(tmp_path, monkeypatch) -> None:
    """migration 未適用（テーブル欠落）は生の psycopg 文字列でなく db_schema_unready コード＋対処を返す（D-1）。

    ユーザー実障害の再現：DATABASE_URL 接続済みだが `classes`(migration 0007) が未整備でクラスを作ると
    record_store は fail-loud（no such table）。それを保育士に意味の通る文言（alembic upgrade head）へ翻訳する。
    """
    monkeypatch.setattr(settings, "database_url", f"sqlite:///{tmp_path}/empty.db")
    record_store.reset_engine_cache()  # スキーマは作らない（create_all を呼ばない）
    r = _client().post("/api/classes", json={"name": "ひまわり", "fiscal_year": ""})
    record_store.reset_engine_cache()
    assert r.status_code == 200  # 既存の {status:error} 契約は維持（HTTP は 200・本文で伝える）
    body = r.json()
    assert body["status"] == "error"
    assert body["code"] == "db_schema_unready"
    assert "alembic upgrade head" in body["detail"]
    assert "no such table" not in body["detail"]  # 生の SQL エラーは露出しない


def test_records_degrade_when_db_unset(monkeypatch) -> None:
    """DATABASE_URL 未設定＝降格（保存 skipped・一覧/児童は空・config は未接続を正直に返す）。"""
    monkeypatch.setattr(settings, "database_url", "")
    record_store.reset_engine_cache()
    c = _client()
    assert c.get("/api/config").json()["records_connected"] is False
    r = c.post(
        "/api/records", json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"}
    )
    assert r.json()["status"] == "skipped"
    assert c.get("/api/children").json() == {"children": [], "store": "disabled"}
    assert c.get("/api/records").json()["documents"] == []


# ──────────── 初回ログインの auto-seed ＋「データを初期化」（/api/account/reset） ────────────


def test_account_reset_requires_signin() -> None:
    """未サインインは 403（fail-closed）＝deletion-request と同じゲート。"""
    assert _client().post("/api/account/reset").status_code == 403


def test_first_signin_triggers_seed_once(records_db, monkeypatch, _stub_auto_seed) -> None:
    """初回サインイン（新規 workspace）で seed が1回だけ発火し、以降のリクエストでは再発火しない。"""
    c = _client()
    _sign_in_with_google(c, monkeypatch)  # /auth/google の provision_user が発火点
    assert len(_stub_auto_seed) == 1
    workspace_id, kwargs = _stub_auto_seed[0]
    user = record_store.touch_user(
        "sensei@example.com", google_subject="google-subject-123", now=datetime.now()
    )
    assert workspace_id == user["workspace_id"]
    # 2回目以降（既存 workspace）は発火しない
    c.get("/api/config")
    c.get("/api/children")
    assert len(_stub_auto_seed) == 1


def test_first_signin_seeds_default_data(records_db, monkeypatch, real_seed) -> None:
    """初回サインインで名簿・クラス2・確定書類チェーンが実際に入る（全タブが初見で埋まる）。"""
    from hoiku_agent.harness import demo_seed_data as seed_data

    c = _client()
    _sign_in_with_google(c, monkeypatch)
    names = {ch["display_name"] for ch in c.get("/api/children").json()["children"]}
    assert set(seed_data.SEEDED_CHILDREN) <= names  # 名簿10人＋卒園児が登録される
    classes = c.get("/api/classes").json()["classes"]
    assert {cls["name"] for cls in classes} == {"ひよこ組", "あおぞら組"}
    docs = c.get("/api/records").json()["documents"]
    assert len(docs) == sum(len(entries) for _, entries in seed_data.JOBS)
    # 承認フロー体感用に一部は未承認（finalized）で残る
    statuses = {d["status"] for d in docs}
    assert statuses == {"approved", "finalized"}


def test_account_reset_restores_seed(records_db, monkeypatch, real_seed) -> None:
    """初期化＝保育士の追加データが消えデフォルト seed に戻る。ログイン（session）は継続。"""
    from hoiku_agent.harness import demo_seed_data as seed_data

    c = _client()
    _sign_in_with_google(c, monkeypatch)
    r = c.post(
        "/api/children",
        json={"given_name": "たろう", "family_name": "テスト", "gender": "male"},
    )
    assert r.json().get("status") in ("created", "exists")
    seeded = len(seed_data.SEEDED_CHILDREN)
    assert len(c.get("/api/children").json()["children"]) == seeded + 1

    reset = c.post("/api/account/reset").json()
    assert reset["status"] == "ok" and reset.get("purged") is True
    names = {ch["display_name"] for ch in c.get("/api/children").json()["children"]}
    assert "たろうくん" not in names  # 追加した児は消える
    assert set(seed_data.SEEDED_CHILDREN) == names  # seed（名簿＋卒園児）に戻る
    # session は生きている（初期化はログアウトしない）＝サインイン必須 API がそのまま通る
    assert c.post("/api/account/reset").json()["status"] == "ok"


def test_records_save_edit_approve_flow(records_db) -> None:
    """AI 確定→編集保存→承認 のアーカイブ一連（版が積まれ・児童が登録され・証跡が残る）。"""
    c = _client()
    entry = _edit_diary_entry()
    # AI 確定（author_kind=ai）
    r1 = c.post(
        "/api/records",
        json={"kind": "diary", "entry": entry, "rendered_text": "整形", "author_kind": "ai"},
    ).json()
    assert r1["status"] == "saved" and r1["version_seq"] == 1
    # 保育士の編集保存（同一書類に版が積まれる）
    r2 = c.post(
        "/api/records",
        json={
            "kind": "diary",
            "entry": entry,
            "author_kind": "caregiver",
            "actor": "保育士A",
        },
    ).json()
    assert r2["status"] == "saved" and r2["version_seq"] == 2
    assert r2["document_id"] == r1["document_id"]
    # 承認（証跡＝actor が残る）
    r3 = c.post(
        "/api/records/approve", json={"kind": "diary", "entry": entry, "actor": "園長"}
    ).json()
    assert r3["status"] == "approved"
    docs = c.get("/api/records").json()
    assert docs["store"] == "ok"
    assert docs["documents"][0]["status"] == "approved"
    # 児童マスタへ auto-create → /api/children で選択肢に出る
    children = c.get("/api/children").json()["children"]
    assert [x["display_name"] for x in children] == ["架空児A"]
    # 監査証跡（誰が・いつ・何を）
    actions = [(e["action"], e["actor"]) for e in record_store.list_audit_events()]
    assert ("approve", "園長") in actions and ("edit", "保育士A") in actions


class _ApprovalMemorySpy:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.calls: list[dict] = []

    async def add_memory(self, **kwargs) -> None:
        self.calls.append(kwargs)
        if self.fail:
            raise RuntimeError("temporary memory failure")


def test_approval_syncs_saved_version_to_memory_once(records_db, monkeypatch) -> None:
    """実Web順序＝保存→承認で同期し、同じ版の再承認は二重投入しない。"""
    memory = _ApprovalMemorySpy()
    monkeypatch.setattr(settings, "agent_engine_id", "test-memory")
    monkeypatch.setattr(server.app.state, "approval_memory_service", memory)
    c = _client()
    entry = _edit_diary_entry()
    saved = c.post(
        "/api/records",
        json={"kind": "diary", "entry": entry, "author_kind": "caregiver"},
    ).json()

    body = c.post(
        "/api/records/approve",
        json={
            "kind": "diary",
            "entry": entry,
            "actor": "園長",
            "expected_version_seq": saved["version_seq"],
        },
    ).json()
    assert body["status"] == "approved" and body["memory_status"] == "synced"
    assert len(memory.calls) == 1
    assert memory.calls[0]["user_id"] == "caregiver"
    assert "child_id=架空児A" in memory.calls[0]["memories"][0].content.parts[0].text

    again = c.post(
        "/api/records/approve",
        json={
            "kind": "diary",
            "entry": entry,
            "actor": "園長",
            "expected_version_seq": saved["version_seq"],
        },
    ).json()
    assert again["memory_status"] == "already_synced"
    assert len(memory.calls) == 1


def test_memory_failure_keeps_document_unapproved(records_db, monkeypatch) -> None:
    """接続済みMemory Bankが失敗したら503にし、承認済みの偽表示を作らない。"""
    memory = _ApprovalMemorySpy(fail=True)
    monkeypatch.setattr(settings, "agent_engine_id", "test-memory")
    monkeypatch.setattr(server.app.state, "approval_memory_service", memory)
    c = _client()
    entry = _edit_diary_entry()
    saved = c.post(
        "/api/records",
        json={"kind": "diary", "entry": entry, "author_kind": "caregiver"},
    ).json()

    response = c.post(
        "/api/records/approve",
        json={
            "kind": "diary",
            "entry": entry,
            "expected_version_seq": saved["version_seq"],
        },
    )
    assert response.status_code == 503
    assert response.json()["code"] == "memory_write_failed"
    assert c.get("/api/records").json()["documents"][0]["status"] == "finalized"


def test_add_child_registers_real_name_and_gender(records_db) -> None:
    """新規児登録＝本名（姓/名）＋性別。呼び名＋敬称（性別導出）＝display_name をサーバが合成する。"""
    c = _client()
    r = c.post(
        "/api/children",
        json={"family_name": "佐藤", "given_name": "はると", "gender": "male"},
    ).json()
    assert r["status"] == "created"
    assert r["display_name"] == "はるとくん"  # 名＋敬称（男→くん）
    assert r["official_name"] == "佐藤　はると"  # 氏名欄用の本名（姓＋名）
    # /api/children の選択肢に本名つきで出る
    row = next(
        x for x in c.get("/api/children").json()["children"] if x["display_name"] == "はるとくん"
    )
    assert row["given_name"] == "はると" and row["gender"] == "male"
    # 女の子は「ちゃん」
    r2 = c.post("/api/children", json={"given_name": "ゆい", "gender": "female"}).json()
    assert r2["display_name"] == "ゆいちゃん"


def test_add_child_validates_input(records_db) -> None:
    c = _client()
    assert c.post("/api/children", json={"given_name": "", "gender": "male"}).status_code == 400
    assert (
        c.post("/api/children", json={"given_name": "そら", "gender": "その他"}).status_code == 400
    )


def test_add_child_degrades_when_db_unset(monkeypatch) -> None:
    """DB 未設定は skipped（永続化なし）だが display_name は合成して返す（フロントはセッション追加）。"""
    monkeypatch.setattr(settings, "database_url", "")
    record_store.reset_engine_cache()
    r = _client().post("/api/children", json={"given_name": "はると", "gender": "male"}).json()
    assert r["status"] == "skipped" and r["display_name"] == "はるとくん"


def test_add_child_stores_birthdate(records_db) -> None:
    """生年月日つきで登録すると児童マスタに保存され /api/children で返る（月齢自動導出の素）。"""
    c = _client()
    r = c.post(
        "/api/children",
        json={"given_name": "はると", "gender": "male", "birthdate": "2024-11-20"},
    ).json()
    assert r["status"] == "created"
    row = next(
        x for x in c.get("/api/children").json()["children"] if x["display_name"] == "はるとくん"
    )
    assert row["birthdate"] == "2024-11-20"


def test_add_child_rejects_bad_birthdate(records_db) -> None:
    """生年月日が ISO 形式でないときは 400（握りつぶさず可視化）。"""
    r = _client().post(
        "/api/children",
        json={"given_name": "はると", "gender": "male", "birthdate": "2024/11/20"},
    )
    assert r.status_code == 400


def test_finalize_edit_child_record_fills_age_months_from_birthdate(records_db) -> None:
    """保育経過記録の歳児欄＝生年月日から満年齢（○歳○か月）を期末時点で自動充填する（登録済みの子）。"""
    c = _client()
    # 生年月日つきで園児を登録（表示名 はるとくん）
    c.post(
        "/api/children",
        json={"given_name": "はると", "gender": "male", "birthdate": "2024-11-20"},
    )
    entry = _edit_child_record_entry()
    entry["child_id"] = "はるとくん"  # マスタに存在＝自動導出の対象
    entry["period"] = "2026-04〜2026-06"  # 期末＝2026-06 末 → 1歳7か月（手入力値を上書き）
    entry["age_months"] = "9歳9か月"  # わざと誤った手入力＝生年月日が権威で上書きされる
    body = c.post("/api/finalize-edit", json={"kind": "child_record", "entry": entry}).json()
    assert body["ok"] is True
    assert "1歳7か月" in body["formatted"]
    assert "9歳9か月" not in body["formatted"]


def test_finalize_edit_child_record_keeps_manual_age_when_no_birthdate(records_db) -> None:
    """生年月日が未登録の子（架空児）は手入力の月齢を温存する（自動導出しない＝§14 の設計を壊さない）。"""
    entry = _edit_child_record_entry()  # child_id=架空児A（マスタ不在・生年月日なし）
    entry["age_months"] = "1歳3か月"
    body = (
        _client().post("/api/finalize-edit", json={"kind": "child_record", "entry": entry}).json()
    )
    assert body["ok"] is True
    assert "1歳3か月" in body["formatted"]


def test_add_child_is_available_to_signed_in_workspace(records_db) -> None:
    """児童マスタは Google ログイン済み workspace 内の通常操作で、別パスコードを求めない。"""
    c = _client()
    assert c.get("/api/children").status_code == 200
    ok = c.post(
        "/api/children",
        json={"given_name": "はると", "gender": "male"},
    )
    assert ok.status_code == 200 and ok.json()["display_name"] == "はるとくん"


# ──────────────────── クラス（組）マスタ（/api/classes＝record_store 中継・名簿管理） ────────────────────


def test_classes_crud_and_roster_flow(records_db) -> None:
    """クラス作成→児童をクラスへ割当→roster/一覧に反映（園の名簿管理・日誌 roster の素）。"""
    c = _client()
    # クラス作成
    r = c.post("/api/classes", json={"name": "ひまわり組", "fiscal_year": "2026"}).json()
    assert r["status"] == "created" and r["name"] == "ひまわり組" and r["age_bands"] == []
    cid = r["id"]
    # 児童登録（クラス指定つき＝1操作で割当まで完結）
    reg = c.post(
        "/api/children",
        json={"given_name": "はると", "gender": "male", "birthdate": "2021-04-02", "class_id": cid},
    ).json()
    assert reg["status"] == "created" and reg.get("assign") == "ok"
    # 児童登録（後から assign）
    c.post(
        "/api/children", json={"given_name": "ゆい", "gender": "female", "birthdate": "2021-04-02"}
    )
    asg = c.post("/api/classes/assign", json={"child": "ゆいちゃん", "class_id": cid}).json()
    assert asg["status"] == "ok" and asg["class_name"] == "ひまわり組"
    # roster＝クラスの在籍児（日誌フォームの素）
    roster = c.get("/api/classes/roster", params={"class_id": cid}).json()["children"]
    assert [x["display_name"] for x in roster] == ["はるとくん", "ゆいちゃん"]
    # 一覧＋在籍児数
    lst = c.get("/api/classes").json()
    assert lst["store"] == "ok"
    cls = next(x for x in lst["classes"] if x["id"] == cid)
    assert cls["child_count"] == 2 and cls["age_bands"] == ["3-5"]
    # 未所属へ戻す
    assert (
        c.post("/api/classes/assign", json={"child": "ゆいちゃん", "class_id": ""}).json()["status"]
        == "ok"
    )
    assert [
        x["display_name"]
        for x in c.get("/api/classes/roster", params={"class_id": cid}).json()["children"]
    ] == ["はるとくん"]


def test_classes_validate_input(records_db) -> None:
    c = _client()
    assert c.post("/api/classes", json={"name": ""}).status_code == 400
    assert c.post("/api/classes", json={"name": "ばら組", "fiscal_year": "2026"}).status_code == 200
    assert c.post("/api/classes/assign", json={"child": "", "class_id": "x"}).status_code == 400


def test_classes_degrade_when_db_unset(monkeypatch) -> None:
    monkeypatch.setattr(settings, "database_url", "")
    record_store.reset_engine_cache()
    c = _client()
    assert c.get("/api/classes").json() == {"classes": [], "store": "disabled"}
    assert c.get("/api/classes/roster", params={"class_id": "x"}).json()["children"] == []
    assert c.post("/api/classes", json={"name": "ひまわり組"}).json()["status"] == "skipped"


def test_classes_write_is_available_to_signed_in_workspace(records_db) -> None:
    """クラスの書込は Google ログイン済み workspace 内の通常操作である。"""
    c = _client()
    assert c.get("/api/classes").status_code == 200
    assert c.get("/api/classes/roster", params={"class_id": "x"}).status_code == 200
    ok = c.post("/api/classes", json={"name": "ひまわり組"})
    assert ok.status_code == 200 and ok.json()["status"] == "created"


# ──────────────────── 校正AI（/api/proofread＝日本語チェック・言い換え提案・LLM 口） ────────────────────


def test_proofread_collect_items_extracts_prose_only() -> None:
    """校正対象は叙述文（プロース）に限る＝数量的な生活記録・タグ・日付・仮名は渡さない（§14）。"""
    from hoiku_agent.web import proofread

    entry = _edit_diary_entry()  # 架空児A・observed_state/evaluation/practice_record を持つ
    items = proofread.collect_items("diary", entry)
    paths = {it["path"] for it in items}
    assert "practice_record" in paths
    assert "individual_notes[0].observed_state" in paths
    assert "evaluation.child_focus" in paths and "evaluation.self_review" in paths
    # 生活記録（食事/睡眠…）・日付・タグ・child_id は校正対象にしない（AI に事実を触らせない）。
    assert not any("life_record" in p for p in paths)
    assert not any(p == "date" or "tags" in p or "child_id" in p for p in paths)
    # ラベルには子どもの呼び名が文脈として付く（個別記録）。
    note_item = next(it for it in items if it["path"] == "individual_notes[0].observed_state")
    assert "架空児A" in note_item["label"]


def test_proofread_empty_narrative_returns_no_suggestions_without_llm() -> None:
    """叙述文が空なら LLM を呼ばず suggestions 空を返す（正常・非課金の早期リターン）。"""
    c = _client()
    entry = {
        "date": "2026-07-06",
        "age_band": "0-2",
        "attendance": [],
        "individual_notes": [],
        "evaluation": {},
    }
    r = c.post("/api/proofread", json={"kind": "diary", "entry": entry})
    assert r.status_code == 200
    body = r.json()
    assert body["suggestions"] == [] and body["checked"] == 0 and body["error"] is None


def test_proofread_requires_google_login_when_enabled(monkeypatch) -> None:
    """校正AI は LLM 口なので Google ログインが必要。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    assert c.post("/api/proofread", json={"kind": "diary", "entry": {}}).status_code == 401


def test_export_pdf_nursery_uses_registered_real_name(records_db) -> None:
    """要録の氏名欄は登録済みの本名で解決できる（get_child 経路が通り PDF を返す）。"""
    c = _client()
    c.post("/api/children", json={"family_name": "佐藤", "given_name": "はると", "gender": "male"})
    entry = {**_edit_nursery_record_entry(), "child_id": "はるとくん"}
    _assert_is_pdf(c.post("/api/export-pdf", json={"kind": "nursery_record", "entry": entry}))


def test_child_record_entries_endpoint_seeds_youroku(records_db) -> None:
    """保育要録 L4 の seed 取得口＝指定児の保育経過記録（最新版・期間順）を返す（読取なので非ゲート）。"""
    c = _client()
    for period, overall in [("2026-04〜2026-06", "1期"), ("2026-07〜2026-09", "2期")]:
        c.post(
            "/api/records",
            json={
                "kind": "child_record",
                "entry": {
                    "period": period,
                    "age_band": "3-5",
                    "child_id": "架空児A",
                    "development_notes": [{"description": overall, "tags": ["健康"]}],
                    "overall_note": overall,
                },
                "author_kind": "ai",
            },
        )
    body = c.get("/api/records/child-record-entries", params={"child": "架空児A"}).json()
    assert body["store"] == "ok"
    assert [e["period"] for e in body["entries"]] == ["2026-04〜2026-06", "2026-07〜2026-09"]
    # exclude_period＝作成対象の期を「前回まで」seed から除く（保育経過記録の自己履歴・依存モデル 2026-07）
    past = c.get(
        "/api/records/child-record-entries",
        params={"child": "架空児A", "exclude_period": "2026-07〜2026-09"},
    ).json()
    assert [e["period"] for e in past["entries"]] == ["2026-04〜2026-06"]
    # 未登録児は空＝フロントがサンプルへ降格
    assert (
        c.get("/api/records/child-record-entries", params={"child": "未登録"}).json()["entries"]
        == []
    )


def test_class_monthly_seed_endpoint_composes_three_inputs(records_db) -> None:
    """クラス月案 seed 取得口＝①クラス児童の経過記録 ②過去クラス月案 ③未反映期間の日誌の合成（非ゲート）。"""
    c = _client()
    # 経過記録（4〜6月をカバー＝境界 6/30）・6月の日誌（反映済み）・7月の日誌（未反映）・7月のクラス月案
    c.post(
        "/api/records",
        json={
            "kind": "child_record",
            "entry": {
                "period": "2026-04〜2026-06",
                "age_band": "0-2",
                "child_id": "架空児A",
                "development_notes": [{"description": "経過", "tags": ["健康"]}],
                "overall_note": "所見",
            },
            "author_kind": "ai",
        },
    )
    for day in ("2026-06-25", "2026-07-10"):
        c.post(
            "/api/records",
            json={
                "kind": "diary",
                "entry": {**_edit_diary_entry(), "date": day},
                "author_kind": "caregiver",
            },
        )
    c.post(
        "/api/records",
        json={
            "kind": "class_monthly",
            "entry": {"month": "2026-07", "age_band": "0-2", "monthly_goal": "7月の目標"},
            "author_kind": "ai",
        },
    )
    body = c.get(
        "/api/records/class-monthly-seed", params={"age_band": "0-2", "month": "2026-08"}
    ).json()
    assert body["store"] == "ok"
    assert [e["date"] for e in body["class_diary_entries"]] == ["2026-07-10"]  # 未反映だけ
    assert [r["period"] for r in body["class_record_entries"]] == ["2026-04〜2026-06"]
    assert [p["month"] for p in body["past_class_plans"]] == ["2026-07"]
    assert body["class_roster"] == []  # 名簿未整備＝空を正直に返す（フロントは「名簿未登録」表示）
    # month 不正は 400（黙って誤解釈しない）
    bad = c.get("/api/records/class-monthly-seed", params={"age_band": "0-2", "month": "abc"})
    assert bad.status_code == 400


def test_records_write_is_available_to_signed_in_workspace(records_db) -> None:
    """アーカイブの書込は Google ログイン済み workspace に紐付く通常操作である。"""
    c = _client()
    payload = {"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"}
    assert c.get("/api/records").status_code == 200
    assert c.get("/api/children").status_code == 200
    ok = c.post("/api/records", json=payload)
    assert ok.status_code == 200 and ok.json()["status"] == "saved"


def test_records_diary_entries_returns_seed(records_db) -> None:
    """seed 取得口＝期間内の日誌 entry（最新版）を返す（月案 L2／保育経過記録 L3 の還流を DB から）。"""
    c = _client()
    c.post(
        "/api/records", json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"}
    )
    body = c.get(
        "/api/records/diary-entries", params={"date_from": "2026-06-01", "date_to": "2026-06-30"}
    ).json()
    assert body["store"] == "ok"
    assert [e["date"] for e in body["entries"]] == ["2026-06-25"]
    # 範囲外は空
    empty = c.get(
        "/api/records/diary-entries", params={"date_from": "2026-07-01", "date_to": "2026-07-31"}
    ).json()
    assert empty["entries"] == []
    # 不正日付は 400（黙って全件を返さない）
    bad = c.get("/api/records/diary-entries", params={"date_from": "abc", "date_to": "2026-07-31"})
    assert bad.status_code == 400


def test_records_diary_meta_flags_unfilled_evaluation(records_db) -> None:
    """diary-meta＝クラス月案作成時の未記入検出用（id/日付/評価充足）。2視点とも記入で complete。"""
    c = _client()
    full = _edit_diary_entry()  # 2026-06-25・評価2視点あり
    unfilled = dict(
        _edit_diary_entry(), date="2026-06-26", evaluation={"child_focus": "x", "self_review": ""}
    )
    for e in (full, unfilled):
        c.post("/api/records", json={"kind": "diary", "entry": e, "author_kind": "ai"})
    body = c.get(
        "/api/records/diary-meta", params={"date_from": "2026-06-01", "date_to": "2026-06-30"}
    ).json()
    assert body["store"] == "ok"
    by_date = {m["date"]: m for m in body["entries"]}
    assert by_date["2026-06-25"]["evaluation_complete"] is True
    assert by_date["2026-06-26"]["evaluation_complete"] is False  # (b) 空＝未記入
    assert all(m["id"] for m in body["entries"])  # id＝飛んで編集する導線に使う
    # 不正日付は 400（黙って全件を返さない）
    assert (
        c.get(
            "/api/records/diary-meta", params={"date_from": "x", "date_to": "2026-06-30"}
        ).status_code
        == 400
    )


def test_get_record_returns_full_document(records_db) -> None:
    """GET /api/records/{id}＝「書類を見る」タブの詳細（現行版の整形テキスト＋本文 entry）。"""
    c = _client()
    saved = c.post(
        "/api/records",
        json={
            "kind": "diary",
            "entry": _edit_diary_entry(),
            "rendered_text": "整形本文",
            "author_kind": "ai",
        },
    ).json()
    r = c.get(f"/api/records/{saved['document_id']}")
    assert r.status_code == 200
    body = r.json()
    assert body["doc_type"] == "diary"
    assert body["rendered_text"] == "整形本文"
    assert body["author_kind"] == "ai"
    assert body["entry"]["date"] == "2026-06-25"


def test_get_record_missing_or_invalid_404(records_db) -> None:
    """不在・不正 id は 404（500 にしない・偽の中身を出さない）。リテラル路より後宣言でも衝突しない。"""
    import uuid

    c = _client()
    assert c.get(f"/api/records/{uuid.uuid4()}").status_code == 404
    assert c.get("/api/records/not-a-uuid").json()["code"] == "not_found"
    # 併存するリテラル路（diary-entries）は {document_id} に飲まれない（宣言順で優先）。
    assert (
        c.get(
            "/api/records/diary-entries",
            params={"date_from": "2026-07-01", "date_to": "2026-07-31"},
        ).status_code
        == 200
    )


def test_feedback_save_and_list_flow(records_db) -> None:
    """👍👎＋ひとことを書類に紐付けて保存し、一覧で返す（version_seq・actor も添う）。"""
    c = _client()
    saved = c.post(
        "/api/records", json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"}
    ).json()
    doc_id = saved["document_id"]
    r = c.post(
        "/api/records/feedback",
        json={
            "document_id": doc_id,
            "verdict": "down",
            "comment": "もう少し具体的に",
            "actor": "保育士A",
        },
    ).json()
    assert r["status"] == "saved" and r["version_seq"] == 1
    # 一覧（GET）はリテラル路として {document_id} に飲まれず 200・保存済みを返す。
    body = c.get("/api/records/feedback", params={"document_id": doc_id}).json()
    assert body["store"] == "ok"
    assert len(body["feedback"]) == 1
    fb = body["feedback"][0]
    assert fb["verdict"] == "down"
    assert fb["comment"] == "もう少し具体的に"
    assert fb["actor"] == "保育士A"
    assert fb["version_seq"] == 1


def test_feedback_write_is_available_to_signed_in_workspace(records_db) -> None:
    """フィードバックも Google ログイン済み workspace 内で保存する。"""
    c = _client()
    doc_id = c.post(
        "/api/records",
        json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"},
    ).json()["document_id"]
    payload = {"document_id": doc_id, "verdict": "up", "comment": "良い"}
    assert c.get("/api/records/feedback", params={"document_id": doc_id}).status_code == 200
    ok = c.post("/api/records/feedback", json=payload)
    assert ok.status_code == 200 and ok.json()["status"] == "saved"


def test_feedback_google_actor_precedence(records_db, monkeypatch) -> None:
    """フィードバックの actor も Google 検証済み email が自己申告に優先する。"""
    c = _client()
    _sign_in_with_google(c, monkeypatch)
    doc_id = c.post(
        "/api/records",
        json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"},
    ).json()["document_id"]
    c.post(
        "/api/records/feedback",
        json={"document_id": doc_id, "verdict": "up", "comment": "x", "actor": "なりすまし名義"},
    )
    body = c.get("/api/records/feedback", params={"document_id": doc_id}).json()
    assert body["feedback"][0]["actor"] == "sensei@example.com"


def test_feedback_degrades_when_db_unset(monkeypatch) -> None:
    """DB 未接続はフィードバック保存 skipped・一覧空（本流を壊さない補助シグナル）。"""
    monkeypatch.setattr(settings, "database_url", "")
    record_store.reset_engine_cache()
    c = _client()
    r = c.post("/api/records/feedback", json={"document_id": "x", "verdict": "up", "comment": "y"})
    assert r.json()["status"] == "skipped"
    assert c.get("/api/records/feedback").json() == {"feedback": [], "store": "disabled"}


def test_get_record_read_is_not_llm_budgeted(records_db) -> None:
    """単一書類の閲覧は LLM を呼ばず、利用枠も消費しない。"""
    c = _client()
    saved = c.post(
        "/api/records",
        json={"kind": "diary", "entry": _edit_diary_entry(), "author_kind": "ai"},
    ).json()
    assert c.get(f"/api/records/{saved['document_id']}").status_code == 200


# ──────────────────── 保育経過記録 年間マトリクスの過去期埋め込み（chohyo_pdf × record_store） ────────────────────


def test_assign_period_columns_fills_same_fiscal_year_only() -> None:
    """期→列割当の純関数＝同じ年度だけ他列へ・今回の期が常に優先・読めない期/別児は除外。"""
    from hoiku_agent.web.chohyo_pdf import assign_period_columns

    current = {"period": "2026-07〜2026-09", "child_id": "はるとくん", "overall_note": "今回"}
    past_q1 = {"period": "2026-04〜2026-06", "child_id": "はるとくん", "overall_note": "1期"}
    other_year = {"period": "2025-04〜2025-06", "child_id": "はるとくん"}  # 前年度＝別シート
    unparseable = {"period": "1学期", "child_id": "はるとくん"}  # 期が読めない＝誤配置しない
    other_child = {"period": "2026-10〜2026-12", "child_id": "めいちゃん"}  # 別児＝防御的に除外
    stale_same_q = {"period": "2026-07〜2026-09", "child_id": "はるとくん", "overall_note": "旧版"}
    cols = assign_period_columns(
        current, [past_q1, other_year, unparseable, other_child, stale_same_q]
    )
    assert set(cols) == {0, 1}
    assert cols[0] is past_q1
    assert cols[1] is current  # アーカイブの旧版より今回の entry が勝つ


def test_assign_period_columns_unparseable_current_stays_alone() -> None:
    """今回の期が読めなければ年度を同定できない＝過去期は埋めず先頭列に今回だけ（誤った列に描かない）。"""
    from hoiku_agent.web.chohyo_pdf import assign_period_columns

    current = {"period": "1学期", "child_id": "はるとくん"}
    past = {"period": "2026-04〜2026-06", "child_id": "はるとくん"}
    assert assign_period_columns(current, [past]) == {0: current}


def test_export_pdf_child_record_embeds_archived_periods(records_db) -> None:
    """保育経過記録の帳票PDF はアーカイブの同児・同年度の過去期で他列を埋めて返す（未接続は従来どおり）。"""
    past = dict(_edit_child_record_entry(), overall_note="1期の所見")
    saved = record_store.save_document(
        "child_record", past, author_kind="ai", now=datetime(2026, 7, 5)
    )
    assert saved["status"] == "saved"
    current = dict(
        _edit_child_record_entry(),
        period="2026-07〜2026-09",
        overall_note="2期の所見",
    )
    r = _client().post("/api/export-pdf", json={"kind": "child_record", "entry": current})
    _assert_is_pdf(r)


# ──────────────────── Google Sign-In identity（Phase 3・検証済み actor と users） ────────────────────


def test_google_signin_protects_app_and_api(monkeypatch) -> None:
    """認証を有効にすると、案内以外は session 無しに公開しない。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    assert c.get("/app/").headers["location"] == "/?next=/app/"
    assert c.get("/api/config").status_code == 401
    welcome = c.get("/")
    assert "accounts.google.com/gsi/client" in welcome.text
    assert "test-client.apps.googleusercontent.com" in welcome.text


def test_google_popup_callback_is_rendered(monkeypatch) -> None:
    """Googleの外部Origin POSTを避けるため、同一Origin popup callback を描画する。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    welcome = _client().get("/")
    assert 'data-callback="handleGoogleCredential"' in welcome.text
    assert 'fetch("/auth/google"' in welcome.text
    assert "data-login_uri=" not in welcome.text


def test_google_verified_user_becomes_actor_and_provisions_user(records_db, monkeypatch) -> None:
    """検証済み session の email/sub が config・証跡・users に反映される。"""
    c = _client()
    _sign_in_with_google(c, monkeypatch)
    assert c.get("/").headers["location"] == "/app/"  # 再訪は案内を挟まず作業へ戻る。
    assert c.get("/api/config").json()["user_email"] == "sensei@example.com"
    payload = {
        "kind": "diary",
        "entry": _edit_diary_entry(),
        "author_kind": "ai",
        "actor": "なりすまし名義",
    }
    assert c.post("/api/records", json=payload).json()["status"] == "saved"
    approve = {"kind": "diary", "entry": _edit_diary_entry(), "actor": "なりすまし名義"}
    assert c.post("/api/records/approve", json=approve).json()["status"] == "approved"
    assert {e["actor"] for e in record_store.list_audit_events()} == {"sensei@example.com"}
    # users へ auto-provision（表示名は後から DB で設定できる）
    with record_store.Session(record_store._engine()) as session:
        users = list(session.scalars(record_store.sa.select(record_store.User)))
    assert [u.email for u in users] == ["sensei@example.com"]
    assert users[0].google_subject == "google-subject-123"


def test_google_callback_rejects_bad_csrf(monkeypatch) -> None:
    """popup callback は案内画面が発行した CSRF cookie と header が違えば拒否する。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    c.get("/")
    r = c.post("/auth/google", json={"credential": "token"}, headers={"X-Login-CSRF": "wrong"})
    assert r.status_code == 400 and r.json()["code"] == "csrf_failed"


def test_google_callback_uses_double_submit_cookie_not_session(monkeypatch) -> None:
    """本番同様に session cookie が復元できなくても専用CSRF cookieで検証できる。"""
    from hoiku_agent.web import auth

    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    monkeypatch.setattr(
        auth,
        "validate_google_credential",
        lambda credential: auth.GoogleUser(
            subject="google-subject-123", email="sensei@example.com"
        ),
    )
    c = _client()
    welcome = c.get("/")
    csrf = re.search(r'"X-Login-CSRF": "([^"]+)"', welcome.text).group(1)
    assert c.cookies.get(auth.LOGIN_CSRF_COOKIE) == auth.login_csrf_cookie_value(csrf)
    c.cookies.delete("session")

    response = c.post(
        "/auth/google",
        json={"credential": "signed-token"},
        headers={"Origin": "http://testserver", "X-Login-CSRF": csrf},
    )

    assert response.status_code == 200
    assert response.json()["redirect"] == "/app/"
    assert auth.LOGIN_CSRF_COOKIE not in c.cookies


def test_google_login_survives_browser_auto_requests(monkeypatch) -> None:
    """favicon の自動取得や案内画面の再描画が挟まっても、表示中のページからログインできる。

    実ブラウザは案内画面の直後に /favicon.ico を自動要求する。これを認証ガードが /?next=… へ
    流すと案内が裏で再描画され、CSRF token の回転で正しいログインが csrf_failed になっていた
    （本番で全ログインが「ログインをもう一度お試しください」で止まる回帰の再現）。
    """
    from hoiku_agent.web import auth

    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    monkeypatch.setattr(
        auth,
        "validate_google_credential",
        lambda credential: auth.GoogleUser(subject="google-subject-123", email="s@example.com"),
    )
    c = _client()
    welcome = c.get("/")
    csrf = re.search(r'"X-Login-CSRF": "([^"]+)"', welcome.text).group(1)
    # ブラウザの自動 favicon 要求＝公開アセットとして返し、ログイン導線（/?next=…）へ流さない。
    favicon = c.get("/favicon.ico", headers={"Sec-Fetch-Mode": "no-cors"})
    assert favicon.status_code == 200
    assert favicon.headers["content-type"] == "image/png"
    # 別タブ等で案内が再描画されても、有効な cookie の token を使い回す（回転させない）。
    rerender = c.get("/")
    assert re.search(r'"X-Login-CSRF": "([^"]+)"', rerender.text).group(1) == csrf
    r = c.post(
        "/auth/google",
        json={"credential": "signed-token"},
        headers={"Origin": "http://testserver", "X-Login-CSRF": csrf},
    )
    assert r.status_code == 200 and r.json()["redirect"] == "/app/"


def test_unauthenticated_subresource_gets_401_not_redirect(monkeypatch) -> None:
    """ブラウザ自動発行のサブリソース要求（Sec-Fetch-Mode≠navigate）は案内へ流さず 401 を返す。

    案内画面へのリダイレクトは post_login_path を汚染するため、画面遷移（navigate・ヘッダ無しの
    旧環境含む）だけを案内へ戻す。
    """
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    sub = c.get("/app/app.js", headers={"Sec-Fetch-Mode": "cors"})
    assert sub.status_code == 401 and sub.json()["code"] == "auth_required"
    nav = c.get("/app/", headers={"Sec-Fetch-Mode": "navigate"})
    assert nav.status_code == 307 and nav.headers["location"] == "/?next=/app/"


def test_google_callback_rejects_cross_origin_post(monkeypatch) -> None:
    """Google callback は同一Origin fetch に限定し、外部OriginのPOSTをADKが拒否する。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    c = _client()
    welcome = c.get("/")
    csrf = re.search(r'"X-Login-CSRF": "([^"]+)"', welcome.text).group(1)
    r = c.post(
        "/auth/google",
        headers={
            "Origin": "https://accounts.google.com",
            "X-Login-CSRF": csrf,
        },
        json={"credential": "invalid"},
    )
    assert r.status_code == 403 and r.text == "Forbidden: origin not allowed"


def test_set_user_display_name_updates_config_and_actor(records_db, monkeypatch) -> None:
    """POST /api/user＝Google の検証済み identity に表示名を紐づけ、証跡へ反映する。"""
    c = _client()
    _sign_in_with_google(c, monkeypatch)
    r = c.post("/api/user", json={"display_name": "そうた先生"})
    assert r.status_code == 200 and r.json()["display_name"] == "そうた先生"
    # config に user_display_name として乗る（フロントが名前を表示できる）。
    cfg = c.get("/api/config").json()
    assert cfg["user_email"] == "sensei@example.com"
    assert cfg["user_display_name"] == "そうた先生"
    # 以後の保存の証跡 actor は「表示名（email）」＝自己申告の名義は無視される。
    payload = {
        "kind": "diary",
        "entry": _edit_diary_entry(),
        "author_kind": "ai",
        "actor": "無視名義",
    }
    assert c.post("/api/records", json=payload).json()["status"] == "saved"
    assert {e["actor"] for e in record_store.list_audit_events()} == {
        "そうた先生（sensei@example.com）"
    }


def test_set_user_display_name_requires_signin(monkeypatch) -> None:
    """未サインイン（検証済み email なし）は 403＝他人の表示名を書けない（fail-closed）。"""
    r = _client().post(
        "/api/user",
        json={"display_name": "なりすまし"},
    )
    assert r.status_code == 403 and r.json()["code"] == "auth_required"


def test_set_user_display_name_requires_no_extra_secret(records_db, monkeypatch) -> None:
    """/api/user はサインイン済みなら別パスコードなしで自分の表示名を保存できる。"""
    c = _client()
    _sign_in_with_google(c, monkeypatch)
    r = c.post(
        "/api/user",
        json={"display_name": "そうた先生"},
    )
    assert r.status_code == 200 and r.json()["status"] == "ok"


# ──────────────────── 表記ルール辞書（ひらがな表記DX＝/api/notation の中継） ────────────────────


def test_notation_get_returns_rules_and_store() -> None:
    """GET /api/notation は表記ルール一覧＋store を返す（読み取りは素通し）。"""
    body = _client().get("/api/notation").json()
    assert "rules" in body and "store" in body
    assert all(
        {"id", "pattern", "replacement", "kind", "enabled"} <= r.keys() for r in body["rules"]
    )


def test_notation_crud_roundtrip(tmp_path, monkeypatch) -> None:
    """追加→編集→削除が harness 経由で回り、更新後の一覧を返す（repo シードは汚さない）。"""
    from hoiku_agent.harness import notation_store as ns

    monkeypatch.setattr(ns, "_NOTATION_PATH", tmp_path / "表記ルール.json")
    c = _client()
    # 追加
    r = c.post(
        "/api/notation", json={"pattern": "出来た", "replacement": "できた", "note": "補助動詞"}
    )
    assert r.status_code == 200 and r.json()["status"] == "ok"
    rules = r.json()["rules"]
    added = next(x for x in rules if x["pattern"] == "出来た")
    assert added["replacement"] == "できた" and added["enabled"] is True
    rid = added["id"]
    # 編集（無効化）
    r = c.patch(f"/api/notation/{rid}", json={"enabled": False, "replacement": "でけた"})
    assert r.json()["status"] == "ok"
    edited = next(x for x in r.json()["rules"] if x["id"] == rid)
    assert edited["enabled"] is False and edited["replacement"] == "でけた"
    # 削除
    r = c.delete(f"/api/notation/{rid}")
    assert r.json()["status"] == "ok"
    assert all(x["id"] != rid for x in r.json()["rules"])


def test_notation_add_duplicate_pattern_rejected(tmp_path, monkeypatch) -> None:
    from hoiku_agent.harness import notation_store as ns

    monkeypatch.setattr(ns, "_NOTATION_PATH", tmp_path / "表記ルール.json")
    c = _client()
    assert (
        c.post("/api/notation", json={"pattern": "子供", "replacement": "子ども"}).status_code
        == 200
    )
    dup = c.post("/api/notation", json={"pattern": "子供", "replacement": "こども"})
    assert dup.status_code == 409 and dup.json()["status"] == "rejected"


def test_notation_invalid_kind_400(tmp_path, monkeypatch) -> None:
    from hoiku_agent.harness import notation_store as ns

    monkeypatch.setattr(ns, "_NOTATION_PATH", tmp_path / "表記ルール.json")
    r = _client().post(
        "/api/notation", json={"pattern": "x", "replacement": "y", "kind": "怪しい種別"}
    )
    assert r.status_code == 400 and r.json()["status"] == "error"


def test_notation_writes_require_no_extra_secret(tmp_path, monkeypatch) -> None:
    """表記ルール編集は Google ログイン済み workspace の通常操作である。"""
    from hoiku_agent.harness import notation_store as ns

    monkeypatch.setattr(ns, "_NOTATION_PATH", tmp_path / "表記ルール.json")
    c = _client()
    assert c.get("/api/notation").status_code == 200
    assert c.post("/api/notation", json={"pattern": "x", "replacement": "y"}).status_code == 200


# ──────────────────── アップロード取込（upload_extract / /api/parse-upload） ────────────────────
# 実 LLM は回さない（creds 不要・決定論）：抽出（決定的）は実物で、書き起こし（agentic）は _run_parser を
# monkeypatch で差し替える。権威的上書き（対象キー/child/age_band）と harness.finalize 中継・降格を担保する。


def _docx_bytes(*lines: str) -> bytes:
    """テスト用の docx を1段落ずつ作ってバイト列で返す。"""
    import io

    from docx import Document

    d = Document()
    for line in lines:
        d.add_paragraph(line)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_upload_extract_docx_xlsx_pdf() -> None:
    """decision論：docx/xlsx はテキスト抽出、pdf は inline_data、未対応/空は ValueError。"""
    from hoiku_agent.web import upload_extract as ux

    ex = ux.extract_upload("n.docx", None, _docx_bytes("保育日誌 5月15日", "はるとくん 完食"))
    assert ex.fmt == "docx" and "はるとくん" in ex.text and ux.to_parts(ex)[0].text

    ex2 = ux.extract_upload("m.xlsx", None, _xlsx_bytes([["対象月", "2026-06"], ["児童", "めい"]]))
    assert ex2.fmt == "xlsx" and "2026-06" in ex2.text

    ex3 = ux.extract_upload("s.pdf", "application/pdf", b"%PDF-1.4 x")
    assert ex3.fmt == "pdf" and ex3.text == "" and ux.to_parts(ex3)[0].inline_data is not None

    with pytest.raises(ValueError):
        ux.extract_upload("old.doc", None, b"bin")  # 旧バイナリ様式は非対応
    with pytest.raises(ValueError):
        ux.extract_upload("empty.pdf", None, b"")  # 空


def test_parse_upload_overrides_keys_and_relays_finalize(monkeypatch) -> None:
    """取込エンドポイント：LLM が種別/対象を取り違えても保育士入力（与件）で権威的に上書きし、
    harness.finalize で整形・検査した結果を返す（保存は後段 /api/records・ここは解析のみ）。"""
    from hoiku_agent.web import upload_parse

    fence = (
        "読み取りました。\n```json\n"
        '{"period":"WRONG","age_band":"0-2","child_id":"別の子",'
        '"development_notes":[{"description":"友だちと関わって遊ぶ姿が増えた","tags":["人間関係"]}],'
        '"overall_note":"関わりが広がった。"}\n```'
    )

    async def _fake(agent, parts):
        return fence

    monkeypatch.setattr(upload_parse, "_run_parser", _fake)
    c = _client()
    r = c.post(
        "/api/parse-upload",
        data={
            "kind": "child_record",
            "target": "2026-04〜2026-06",
            "child": "はるとくん",
            "age_band": "3-5",
        },
        files={
            "file": ("j.docx", _docx_bytes("保育経過記録 はると 1期"), "application/octet-stream")
        },
    )
    assert r.status_code == 200
    body = r.json()
    # 与件で権威的に上書きされている（LLM の取り違えを封じる）。
    assert body["entry"]["period"] == "2026-04〜2026-06"
    assert body["entry"]["child_id"] == "はるとくん"
    assert body["entry"]["age_band"] == "3-5"
    # harness の整形・検査結果が中継されている。
    assert body["ok"] is True and body["parse_error"] is None
    assert "発達の経過" in body["formatted"]


def test_parse_upload_class_monthly_overrides_and_relays(monkeypatch) -> None:
    """クラス月案の取込（§18）：クラス単位＝主対象児なしで、対象月＝与件を権威的に上書きし、
    grid が正準7行に整えられた harness 確定結果を中継する（kind=class_monthly・class_monthly も取込対応）。"""
    import json as _json

    from hoiku_agent.web import upload_parse

    parsed = {
        "month": "WRONG",  # LLM の取り違え → 与件で上書きされることを確かめる
        "age_band": "0-2",
        "class_name": "ひよこ組",
        "monthly_goal": "梅雨期も健康に過ごす",
        "prev_month_state": "感触遊びに集中していた",
        "grid": [
            {
                "category": c,
                "domain": d,
                "aim": f"{d}のねらい",
                "environment": "環境",
                "child_state": "姿",
                "support": "配慮",
            }
            for c, d in (
                ("養護", "生命の保持"),
                ("養護", "情緒の安定"),
                ("教育", "健康"),
                ("教育", "人間関係"),
                ("教育", "環境"),
                ("教育", "言葉"),
                ("教育", "表現"),
            )
        ],
        "individual_goals": [
            {
                "child_id": "はるとくん",
                "age_months": "1歳3か月",
                "child_state": "歩行が安定",
                "aim_support": "探索を見守る",
            }
        ],
    }
    fence = "読み取りました。\n```json\n" + _json.dumps(parsed, ensure_ascii=False) + "\n```"

    async def _fake(agent, parts):
        return fence

    monkeypatch.setattr(upload_parse, "_run_parser", _fake)
    r = _client().post(
        "/api/parse-upload",
        data={"kind": "class_monthly", "target": "2026-07", "child": "", "age_band": "0-2"},
        files={
            "file": (
                "m.docx",
                _docx_bytes("月間指導計画 7月 0歳児クラス"),
                "application/octet-stream",
            )
        },
    )
    assert r.status_code == 200
    body = r.json()
    # 対象月は与件で権威的に上書き。クラス単位なので top-level child_id は付かない。
    assert body["entry"]["month"] == "2026-07"
    assert body["entry"]["age_band"] == "0-2"
    assert "child_id" not in body["entry"]
    # grid は正準7行にそろい、0–2 の個人目標は原本のまま残る。
    assert len(body["entry"]["grid"]) == 7
    assert body["entry"]["individual_goals"][0]["child_id"] == "はるとくん"
    # harness の整形・検査結果が中継され、型が成立している。
    assert body["ok"] is True and body["parse_error"] is None
    assert body["formatted"]


def test_parse_upload_degrades_on_llm_failure(monkeypatch) -> None:
    """LLM 失敗（creds 未設定等）は 200＋parse_error で正直に降格し、与件入りの最小 entry を返す。"""
    from hoiku_agent.web import upload_parse

    async def _boom(agent, parts):
        raise RuntimeError("no creds")

    monkeypatch.setattr(upload_parse, "_run_parser", _boom)
    r = _client().post(
        "/api/parse-upload",
        data={"kind": "diary", "target": "2026-05-15", "age_band": "0-2"},
        files={"file": ("n.docx", _docx_bytes("保育日誌"), "application/octet-stream")},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is False and body["parse_error"]  # 正直に降格
    assert body["entry"]["date"] == "2026-05-15"  # 与件は入っている＝手入力で補える


def test_parse_upload_unsupported_format_400(monkeypatch) -> None:
    r = _client().post(
        "/api/parse-upload",
        data={"kind": "diary", "target": "2026-05-15", "age_band": "0-2"},
        files={"file": ("old.doc", b"legacy-binary", "application/msword")},
    )
    assert r.status_code == 400 and r.json()["code"] == "invalid_request"


def test_parse_upload_requires_google_login_when_enabled(monkeypatch) -> None:
    """アップロード取込は LLM 口なので Google ログインが必要。"""
    monkeypatch.setattr(
        settings, "google_oauth_client_id", "test-client.apps.googleusercontent.com"
    )
    monkeypatch.setattr(settings, "session_secret", "test-session-secret")
    r = _client().post(
        "/api/parse-upload",
        data={"kind": "diary", "target": "2026-05-15", "age_band": "0-2"},
        files={"file": ("n.docx", _docx_bytes("x"), "application/octet-stream")},
    )
    assert r.status_code == 401
